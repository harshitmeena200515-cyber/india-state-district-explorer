from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
DARK = RGBColor(15, 23, 42)
CARD = RGBColor(30, 41, 59)
ACCENT = RGBColor(99, 102, 241)
PURPLE = RGBColor(168, 85, 247)
WHITE = RGBColor(248, 250, 252)
GRAY = RGBColor(148, 163, 184)
GREEN = RGBColor(52, 211, 153)
ORANGE = RGBColor(251, 146, 60)
CYAN = RGBColor(34, 211, 238)


def dark_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK

def accent_bar(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.07))
    s.fill.solid()
    s.fill.fore_color.rgb = ACCENT
    s.line.fill.background()

def side_bar(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), prs.slide_height)
    s.fill.solid()
    s.fill.fore_color.rgb = ACCENT
    s.line.fill.background()

def slide_title(slide, text, sub=""):
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.8))
    p = tx.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    # underline
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(2.5), Inches(0.04))
    s.fill.solid()
    s.fill.fore_color.rgb = ACCENT
    s.line.fill.background()
    if sub:
        tx2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(10), Inches(0.4))
        p2 = tx2.text_frame.paragraphs[0]
        p2.text = sub
        p2.font.size = Pt(15)
        p2.font.color.rgb = GRAY

def card(slide, left, top, w, h, lines, color=ACCENT):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = CARD
    s.line.color.rgb = RGBColor(51, 65, 85)
    s.line.width = Pt(1)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.2)
    for i, (txt, sz, clr, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(sz)
        p.font.color.rgb = clr
        p.font.bold = bold
        p.space_after = Pt(4)


# ═══════════════════════════════════════════════
# SLIDE 1 — TITLE
# ═══════════════════════════════════════════════
s1 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s1)
accent_bar(s1)

# decorative circles
for (l, t, sz, c) in [(Inches(10), Inches(-1.5), Inches(5), RGBColor(30, 27, 75)),
                       (Inches(-2), Inches(4.5), Inches(6), RGBColor(30, 27, 75))]:
    o = s1.shapes.add_shape(MSO_SHAPE.OVAL, l, t, sz, sz)
    o.fill.solid()
    o.fill.fore_color.rgb = c
    o.line.fill.background()

tx = s1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10), Inches(1.2))
p = tx.text_frame.paragraphs[0]
p.text = "State & District Finder"
p.font.size = Pt(52)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

tx2 = s1.shapes.add_textbox(Inches(2), Inches(3.2), Inches(9), Inches(0.6))
p2 = tx2.text_frame.paragraphs[0]
p2.text = "A Simple Web Tool to Search Districts of Any Indian State"
p2.font.size = Pt(20)
p2.font.color.rgb = GRAY
p2.alignment = PP_ALIGN.CENTER

# divider
ds = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(4.2), Inches(4), Inches(0.04))
ds.fill.solid()
ds.fill.fore_color.rgb = ACCENT
ds.line.fill.background()

tx3 = s1.shapes.add_textbox(Inches(2), Inches(4.7), Inches(9), Inches(1.5))
tf3 = tx3.text_frame
p3 = tf3.paragraphs[0]
p3.text = "Presented By"
p3.font.size = Pt(13)
p3.font.color.rgb = GRAY
p3.alignment = PP_ALIGN.CENTER

p4 = tf3.add_paragraph()
p4.text = "Harshit  &  Golu"
p4.font.size = Pt(28)
p4.font.bold = True
p4.font.color.rgb = ACCENT
p4.alignment = PP_ALIGN.CENTER

p5 = tf3.add_paragraph()
p5.text = "Maa Bharti PG College"
p5.font.size = Pt(14)
p5.font.color.rgb = GRAY
p5.alignment = PP_ALIGN.CENTER


# ═══════════════════════════════════════════════
# SLIDE 2 — WHAT IS THIS PROJECT?
# ═══════════════════════════════════════════════
s2 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s2)
side_bar(s2)
slide_title(s2, "What is This Project?", "Yeh project kya kaam karta hai")

card(s2, Inches(0.8), Inches(2.2), Inches(11.5), Inches(4.5), [
    ("State & District Finder", 22, ACCENT, True),
    ("", 8, WHITE, False),
    ("Yeh ek simple web application hai jisme aap kisi bhi Indian state", 16, WHITE, False),
    ("ka naam search karte hain aur uske saare districts turant dikhai dete hain.", 16, WHITE, False),
    ("", 8, WHITE, False),
    ("Example:", 16, ORANGE, True),
    ("Agar aap search box mein 'Maharashtra' likhte hain", 15, WHITE, False),
    ("toh screen par Ahmednagar, Akola, Amravati... saare districts aa jayenge.", 15, WHITE, False),
    ("", 8, WHITE, False),
    ("Yeh project bina internet ke bhi kaam karta hai kyunki sara data", 15, GRAY, False),
    ("locally (JSON file mein) stored hai.", 15, GRAY, False),
])


# ═══════════════════════════════════════════════
# SLIDE 3 — PROJECT FILES
# ═══════════════════════════════════════════════
s3 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s3)
side_bar(s3)
slide_title(s3, "Project Files", "Is project mein kaun kaun si files hain")

files = [
    ("index.html", "Main web page — Design + Search logic", "Yeh file browser mein khulti hai", ACCENT),
    ("states-and-districts.json", "Data file — Saare states aur districts ka data", "Yeh file data store karti hai", GREEN),
    ("server.py", "Python server — Website ko locally host karta hai", "Yeh file server chalati hai", ORANGE),
    ("data.py", "Terminal script — Bina browser ke search karo", "Yeh file terminal mein chalti hai", CYAN),
]

for i, (name, desc, role, clr) in enumerate(files):
    col = i % 2
    row = i // 2
    left = Inches(0.8) + col * Inches(6.2)
    top = Inches(2.2) + row * Inches(2.4)
    card(s3, left, top, Inches(5.8), Inches(2.1), [
        (name, 20, clr, True),
        ("", 6, WHITE, False),
        (desc, 15, WHITE, False),
        (role, 13, GRAY, False),
    ])


# ═══════════════════════════════════════════════
# SLIDE 4 — WHAT CODE IS USED & WHY? (Part 1 - HTML)
# ═══════════════════════════════════════════════
s4 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s4)
side_bar(s4)
slide_title(s4, "Code Mein Kya Use Hua? (Part 1)", "HTML & CSS — Design ke liye")

card(s4, Inches(0.8), Inches(2.2), Inches(5.8), Inches(4.5), [
    ("HTML (HyperText Markup Language)", 18, ACCENT, True),
    ("", 6, WHITE, False),
    ("Kya kaam karta hai?", 15, ORANGE, True),
    ("Web page ka structure banata hai — jaise search box,", 14, WHITE, False),
    ("button, heading, aur results area.", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Kyun use kiya?", 15, GREEN, True),
    ("Browser sirf HTML samajhta hai. Bina HTML ke", 14, WHITE, False),
    ("koi bhi web page nahi ban sakta. Yeh har", 14, WHITE, False),
    ("website ki buniyad (foundation) hai.", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Project mein kahan use hua?", 15, CYAN, True),
    ("index.html file mein — search box, button, aur", 14, GRAY, False),
    ("results dikhane ke liye.", 14, GRAY, False),
])

card(s4, Inches(6.9), Inches(2.2), Inches(5.8), Inches(4.5), [
    ("CSS (Cascading Style Sheets)", 18, PURPLE, True),
    ("", 6, WHITE, False),
    ("Kya kaam karta hai?", 15, ORANGE, True),
    ("Web page ko sundar banata hai — colors, fonts,", 14, WHITE, False),
    ("spacing, animations sab CSS se hota hai.", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Kyun use kiya?", 15, GREEN, True),
    ("Bina CSS ke page boring aur plain dikhta hai.", 14, WHITE, False),
    ("CSS se dark theme, glow effects, hover", 14, WHITE, False),
    ("animations aur rounded cards banayi hain.", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Project mein kahan use hua?", 15, CYAN, True),
    ("index.html ke andar <style> tag mein —", 14, GRAY, False),
    ("background, button design, search box styling.", 14, GRAY, False),
])


# ═══════════════════════════════════════════════
# SLIDE 5 — WHAT CODE IS USED & WHY? (Part 2 - JS)
# ═══════════════════════════════════════════════
s5 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s5)
side_bar(s5)
slide_title(s5, "Code Mein Kya Use Hua? (Part 2)", "JavaScript — Search Logic ke liye")

card(s5, Inches(0.8), Inches(2.2), Inches(11.8), Inches(2.0), [
    ("JavaScript (JS)", 20, ACCENT, True),
    ("", 6, WHITE, False),
    ("Kya kaam karta hai?  Web page ko interactive banata hai. Jab user button click kare ya kuch type kare, tab page react kare.", 14, WHITE, False),
    ("Kyun use kiya?  HTML sirf dikhata hai, CSS sirf sajata hai — lekin kaam JavaScript karta hai. Search, data load, results dikhana — sab JS se hota hai.", 14, WHITE, False),
])

# JS code example
code_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.5), Inches(5.8), Inches(2.5))
code_box.fill.solid()
code_box.fill.fore_color.rgb = RGBColor(17, 24, 39)
code_box.line.color.rgb = RGBColor(51, 65, 85)
tf = code_box.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.3)
tf.margin_top = Inches(0.2)

code_title = s5.shapes.add_textbox(Inches(0.8), Inches(4.15), Inches(5), Inches(0.35))
code_title.text_frame.paragraphs[0].text = "Code Example:"
code_title.text_frame.paragraphs[0].font.size = Pt(14)
code_title.text_frame.paragraphs[0].font.bold = True
code_title.text_frame.paragraphs[0].font.color.rgb = ORANGE

js_lines = [
    ("// JSON file se data load karna", GRAY),
    ("fetch('states-and-districts.json')", WHITE),
    ("  .then(res => res.json())", WHITE),
    ("  .then(data => {", WHITE),
    ("    statesData = data.states;", RGBColor(190, 242, 100)),
    ("  });", WHITE),
    ("", WHITE),
    ("// User ki search handle karna", GRAY),
    ("const found = statesData.find(", WHITE),
    ("  s => s.state.toLowerCase() === query", RGBColor(190, 242, 100)),
    (");", WHITE),
]

for i, (line, clr) in enumerate(js_lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = line
    p.font.size = Pt(12)
    p.font.name = "Consolas"
    p.font.color.rgb = clr
    p.space_after = Pt(1)

card(s5, Inches(6.9), Inches(4.5), Inches(5.8), Inches(2.5), [
    ("JS Project Mein Kya Karta Hai?", 16, ACCENT, True),
    ("", 6, WHITE, False),
    ("1. fetch() se JSON file ka data load karta hai", 14, WHITE, False),
    ("2. User jo state name type kare usse match karta hai", 14, WHITE, False),
    ("3. Match milne par districts ki list banata hai", 14, WHITE, False),
    ("4. Har district ko animation ke saath dikhata hai", 14, WHITE, False),
    ("5. Agar state na mile toh error message dikhata hai", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Yeh sab bina page reload kiye hota hai!", 13, GREEN, True),
])


# ═══════════════════════════════════════════════
# SLIDE 6 — WHAT CODE IS USED & WHY? (Part 3 - Python)
# ═══════════════════════════════════════════════
s6 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s6)
side_bar(s6)
slide_title(s6, "Code Mein Kya Use Hua? (Part 3)", "Python — Server & Terminal ke liye")

card(s6, Inches(0.8), Inches(2.2), Inches(5.8), Inches(4.8), [
    ("server.py  (Python)", 20, ORANGE, True),
    ("", 6, WHITE, False),
    ("Kya kaam karta hai?", 15, ACCENT, True),
    ("Ek local web server banata hai taaki", 14, WHITE, False),
    ("index.html browser mein sahi se chale.", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Kyun zaroori hai?", 15, GREEN, True),
    ("Agar hum index.html ko seedha double-click", 14, WHITE, False),
    ("karke kholein, toh browser JSON file ko load", 14, WHITE, False),
    ("nahi karne deta (security reason). Isliye", 14, WHITE, False),
    ("humein ek server chahiye jo files serve kare.", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Python Modules Used:", 15, CYAN, True),
    ("http.server — web pages serve karne ke liye", 13, GRAY, False),
    ("socketserver — network connection handle karne ke liye", 13, GRAY, False),
])

card(s6, Inches(6.9), Inches(2.2), Inches(5.8), Inches(4.8), [
    ("data.py  (Python)", 20, CYAN, True),
    ("", 6, WHITE, False),
    ("Kya kaam karta hai?", 15, ACCENT, True),
    ("Terminal/Command Prompt mein state search", 14, WHITE, False),
    ("karne deta hai — bina browser khole.", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Kyun banaya?", 15, GREEN, True),
    ("Kabhi kabhi browser kholna zaroori nahi hota.", 14, WHITE, False),
    ("Developer ya quick user terminal se seedha", 14, WHITE, False),
    ("state ka naam type karke districts dekh sakte hain.", 14, WHITE, False),
    ("", 6, WHITE, False),
    ("Python Modules Used:", 15, ORANGE, True),
    ("json — JSON file ko read karne ke liye", 13, GRAY, False),
    ("input() — User se state naam lene ke liye", 13, GRAY, False),
    ("for loop — Districts print karne ke liye", 13, GRAY, False),
])


# ═══════════════════════════════════════════════
# SLIDE 7 — WHY JSON?
# ═══════════════════════════════════════════════
s7 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s7)
side_bar(s7)
slide_title(s7, "JSON Data File — Kyun Use Kiya?", "states-and-districts.json ka role")

card(s7, Inches(0.8), Inches(2.2), Inches(5.8), Inches(2.0), [
    ("JSON (JavaScript Object Notation)", 18, GREEN, True),
    ("", 6, WHITE, False),
    ("Yeh ek lightweight data format hai jo human-readable hai.", 14, WHITE, False),
    ("Is project mein saare states aur districts ka data", 14, WHITE, False),
    ("isi file mein stored hai — total 13 KB ka data.", 14, WHITE, False),
])

card(s7, Inches(6.9), Inches(2.2), Inches(5.8), Inches(2.0), [
    ("Database kyun nahi use kiya?", 18, ORANGE, True),
    ("", 6, WHITE, False),
    ("Is project mein data chota aur fixed hai (states change nahi", 14, WHITE, False),
    ("hote). Isliye MySQL ya MongoDB jaisa heavy database lagana", 14, WHITE, False),
    ("zaroori nahi tha. JSON simple aur fast hai.", 14, WHITE, False),
])

# JSON example
code_box2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.6), Inches(11.8), Inches(2.5))
code_box2.fill.solid()
code_box2.fill.fore_color.rgb = RGBColor(17, 24, 39)
code_box2.line.color.rgb = RGBColor(51, 65, 85)
tf2 = code_box2.text_frame
tf2.word_wrap = True
tf2.margin_left = Inches(0.3)
tf2.margin_top = Inches(0.2)

json_title = s7.shapes.add_textbox(Inches(0.8), Inches(4.25), Inches(5), Inches(0.35))
json_title.text_frame.paragraphs[0].text = "JSON Data ka Example:"
json_title.text_frame.paragraphs[0].font.size = Pt(14)
json_title.text_frame.paragraphs[0].font.bold = True
json_title.text_frame.paragraphs[0].font.color.rgb = GREEN

json_lines = [
    ('{ "states": [', WHITE),
    ('    {', WHITE),
    ('      "state": "Goa",', RGBColor(190, 242, 100)),
    ('      "districts": [', WHITE),
    ('          "North Goa",', RGBColor(125, 211, 252)),
    ('          "South Goa"', RGBColor(125, 211, 252)),
    ('      ]', WHITE),
    ('    },', WHITE),
    ('    { "state": "Maharashtra", "districts": ["Ahmednagar", "Akola", ...] }', GRAY),
    ('  ]', WHITE),
    ('}', WHITE),
]

for i, (line, clr) in enumerate(json_lines):
    p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    p.text = line
    p.font.size = Pt(13)
    p.font.name = "Consolas"
    p.font.color.rgb = clr
    p.space_after = Pt(1)


# ═══════════════════════════════════════════════
# SLIDE 8 — HOW IT WORKS (Simple Flow)
# ═══════════════════════════════════════════════
s8 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s8)
side_bar(s8)
slide_title(s8, "Kaise Kaam Karta Hai?", "Step by step samjho")

steps = [
    ("Step 1", "Server Start Karo", "Terminal mein likho:\npython server.py", ACCENT),
    ("Step 2", "Browser Kholo", "Address bar mein likho:\nlocalhost:5000", PURPLE),
    ("Step 3", "State Naam Likho", "Search box mein state\nka naam type karo", GREEN),
    ("Step 4", "Results Dekho", "Screen par saare districts\nlist mein dikhai denge", ORANGE),
]

for i, (num, title, desc, clr) in enumerate(steps):
    left = Inches(0.6) + i * Inches(3.2)

    # Step badge
    badge = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.8), Inches(2.2), Inches(1.3), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = clr
    badge.line.fill.background()
    badge.text_frame.paragraphs[0].text = num
    badge.text_frame.paragraphs[0].font.size = Pt(14)
    badge.text_frame.paragraphs[0].font.bold = True
    badge.text_frame.paragraphs[0].font.color.rgb = WHITE
    badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Arrow
    if i < 3:
        arrow = s8.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left + Inches(2.85), Inches(3.5), Inches(0.5), Inches(0.3))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = GRAY
        arrow.line.fill.background()

    card(s8, left, Inches(2.8), Inches(2.9), Inches(2.8), [
        (title, 18, clr, True),
        ("", 6, WHITE, False),
        (desc, 14, WHITE, False),
    ])

# Bottom summary
card(s8, Inches(0.8), Inches(6.0), Inches(11.8), Inches(1.0), [
    ("Yaad rakho: server.py chalao -> browser mein localhost:5000 kholo -> search karo -> results dekho!", 15, GREEN, True),
])


# ═══════════════════════════════════════════════
# SLIDE 9 — THANK YOU
# ═══════════════════════════════════════════════
s9 = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s9)
accent_bar(s9)

for (l, t, sz, c) in [(Inches(9.5), Inches(-2), Inches(6), RGBColor(30, 27, 75)),
                       (Inches(-2.5), Inches(4), Inches(7), RGBColor(30, 27, 75))]:
    o = s9.shapes.add_shape(MSO_SHAPE.OVAL, l, t, sz, sz)
    o.fill.solid()
    o.fill.fore_color.rgb = c
    o.line.fill.background()

tx = s9.shapes.add_textbox(Inches(2), Inches(2.0), Inches(9), Inches(1.2))
p = tx.text_frame.paragraphs[0]
p.text = "Thank You!"
p.font.size = Pt(58)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

tx2 = s9.shapes.add_textbox(Inches(2), Inches(3.5), Inches(9), Inches(0.6))
p2 = tx2.text_frame.paragraphs[0]
p2.text = "Any Questions?"
p2.font.size = Pt(22)
p2.font.color.rgb = GRAY
p2.alignment = PP_ALIGN.CENTER

ds = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(4.4), Inches(3), Inches(0.04))
ds.fill.solid()
ds.fill.fore_color.rgb = ACCENT
ds.line.fill.background()

tx3 = s9.shapes.add_textbox(Inches(2), Inches(4.9), Inches(9), Inches(1.5))
tf3 = tx3.text_frame
p3 = tf3.paragraphs[0]
p3.text = "Harshit  &  Golu"
p3.font.size = Pt(24)
p3.font.bold = True
p3.font.color.rgb = ACCENT
p3.alignment = PP_ALIGN.CENTER
p4 = tf3.add_paragraph()
p4.text = "Maa Bharti PG College"
p4.font.size = Pt(14)
p4.font.color.rgb = GRAY
p4.alignment = PP_ALIGN.CENTER


# ─── SAVE ─────────────────────────────────
output = r"D:\project\State_District_Finder_PPT.pptx"
prs.save(output)
print(f"\nPPT saved successfully!")
print(f"Location: {output}")
print(f"Total Slides: {len(prs.slides)}")
